"""Invoice spreadsheets: map whatever headers a subsidiary uses onto one schema."""

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from mock491.extract.common import (
    CURRENCY_ALIASES,
    ExtractionError,
    confidence,
    fold,
    parse_tr_number,
    squash,
    vkn_issues,
)
from mock491.schemas import DocType, Extraction, Financials, Lineage

BASE_CONFIDENCE = 0.99  # structured cells: the only real doubt is the content itself

# canonical field -> header spellings seen in the wild (compared after fold())
HEADER_ALIASES = {
    "invoice_no": {"fatura no", "fatura numarası", "belge no"},
    "date": {"fatura tarihi", "tarih"},
    "name": {"firma", "tedarikçi", "tedarikci unvani", "tedarikçi unvanı", "unvan"},
    "vkn": {"vergi no", "vkn", "vergi kimlik no"},
    "amount": {"tutar", "toplam", "fatura tutarı"},
    "currency": {"para birimi", "döviz"},
    "status": {"durum", "odeme durumu", "ödeme durumu"},
}
REQUIRED = {"invoice_no", "name", "amount"}

_AMOUNT_WITH_CURRENCY = re.compile(r"^([\d.,]+)\s*(\S+)?$")


def map_headers(headers: list[object]) -> dict[str, int]:
    lookup = {fold(a): field for field, aliases in HEADER_ALIASES.items() for a in aliases}
    mapping = {}
    for idx, h in enumerate(headers):
        key = fold(str(h or ""))
        if key in lookup:
            mapping[lookup[key]] = idx
    if missing := REQUIRED - mapping.keys():
        raise ExtractionError(f"unrecognised invoice layout, missing columns: {sorted(missing)}")
    return mapping


def _amount_and_currency(cell: object, currency_cell: object) -> tuple[float, str | None]:
    currency = str(currency_cell).strip().upper() if currency_cell else None
    if isinstance(cell, int | float):
        return float(cell), currency
    m = _AMOUNT_WITH_CURRENCY.match(squash(str(cell)))
    if not m:
        raise ExtractionError(f"unparseable amount: {cell!r}")
    return parse_tr_number(m[1]), currency or (m[2] and m[2].upper())


def _date(cell: object) -> date | None:
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, str) and re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", cell.strip()):
        return datetime.strptime(cell.strip(), "%d.%m.%Y").date()
    return None


def extract(path: Path, sha256: str, tenant_id: str) -> list[Extraction]:
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    cols = map_headers(list(next(rows)))
    get = lambda row, field: row[cols[field]] if field in cols else None  # noqa: E731

    out = []
    for row_number, row in enumerate(rows, start=2):
        if not any(row):
            continue
        issues = []
        amount, cur = _amount_and_currency(get(row, "amount"), get(row, "currency"))
        currency = CURRENCY_ALIASES.get(cur or "")
        if currency is None:
            currency, issues = "TRY", issues + ["currency_assumed"]
        vkn_text = str(get(row, "vkn") or "").strip() or None
        issues += vkn_issues(vkn_text)
        doc_date = _date(get(row, "date"))
        if doc_date is None:
            issues.append("date_missing")
        status = get(row, "status")

        out.append(
            Extraction(
                extraction_id=f"EXT-{sha256[:12]}-r{row_number}",
                tenant_id=tenant_id,
                doc_type=DocType.INVOICE,
                document_ref=str(get(row, "invoice_no")),
                document_date=doc_date,
                counterparty_name=squash(str(get(row, "name"))),
                counterparty_vkn=vkn_text if vkn_text and re.fullmatch(r"\d{10}", vkn_text)
                else None,
                financials=Financials(amount=amount, currency=currency),
                paid=fold(str(status)) == "odendi" if status else None,
                lineage=Lineage(sha256=sha256, row_number=row_number,
                                confidence=confidence(BASE_CONFIDENCE, issues), issues=issues),
            )
        )  # fmt: skip
    return out
