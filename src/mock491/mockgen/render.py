"""Turn specs into the files subsidiaries would actually send: contract PDFs, invoice workbooks."""

import os
import random
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook

from mock491 import config

# fpdf2's built-in fonts are Latin-1 only and cannot render ş, ğ, ı, İ.
FONT_CANDIDATES = [
    (config.ROOT / "assets/fonts/DejaVuSans.ttf", config.ROOT / "assets/fonts/DejaVuSans-Bold.ttf"),
    (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
    (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
     Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    (Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
     Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")),
]  # fmt: skip

FILLER_CLAUSES = [
    ("TESLİMAT", "Satıcı, sipariş konusu ürünleri Alıcı'nın bildireceği adrese, sipariş "
     "tarihinden itibaren en geç otuz (30) gün içinde eksiksiz ve ayıpsız olarak teslim etmekle "
     "yükümlüdür. Teslimat sırasında düzenlenen irsaliye ve teslim tutanağı, teslimin ispatında "
     "esas alınır. Kısmi teslimat ancak Alıcı'nın yazılı onayı ile mümkündür."),
    ("KALİTE VE MUAYENE", "Alıcı, teslim edilen ürünleri teslim tarihinden itibaren on beş (15) "
     "gün içinde muayene etme hakkına sahiptir. Muayene sonucunda ayıplı olduğu tespit edilen "
     "ürünler Satıcı tarafından masrafları kendisine ait olmak üzere derhal değiştirilir. Gizli "
     "ayıplara ilişkin ihbar süresi, ayıbın ortaya çıktığı tarihten itibaren başlar."),
    ("GİZLİLİK", "Taraflar, işbu sözleşmenin müzakeresi ve ifası sırasında birbirlerine ait "
     "öğrendikleri ticari sır niteliğindeki her türlü bilgiyi gizli tutmayı, üçüncü kişilerle "
     "paylaşmamayı ve yalnızca sözleşmenin ifası amacıyla kullanmayı kabul eder. Bu yükümlülük "
     "sözleşmenin sona ermesinden sonra beş (5) yıl süreyle devam eder."),
    ("KİŞİSEL VERİLERİN KORUNMASI", "Taraflar, sözleşmenin ifası kapsamında işlenen kişisel "
     "verilerin 6698 sayılı Kişisel Verilerin Korunması Kanunu'na uygun olarak işlenmesi, "
     "saklanması ve imha edilmesi için gerekli her türlü teknik ve idari tedbiri almayı taahhüt "
     "eder."),
    ("MÜCBİR SEBEPLER", "Doğal afet, salgın hastalık, savaş, grev, lokavt ve resmi makamların "
     "kararları gibi tarafların kontrolü dışında gelişen haller mücbir sebep sayılır. Mücbir "
     "sebebin devamı süresince tarafların yükümlülükleri askıya alınır. Mücbir sebebin doksan "
     "(90) günden fazla sürmesi halinde taraflardan her biri sözleşmeyi feshedebilir."),
    ("FESİH", "Taraflardan birinin sözleşmeden doğan yükümlülüklerini ihlal etmesi ve bu ihlali "
     "yazılı ihtara rağmen on (10) gün içinde gidermemesi halinde, diğer taraf sözleşmeyi haklı "
     "sebeple derhal feshedebilir. Fesih, feshe kadar doğmuş hak ve alacakları etkilemez."),
    ("DEVİR YASAĞI", "Satıcı, Alıcı'nın önceden yazılı onayı olmaksızın işbu sözleşmeden doğan "
     "hak ve yükümlülüklerini kısmen veya tamamen üçüncü kişilere devredemez, temlik edemez."),
    ("SİGORTA", "Satıcı, teslimat tamamlanıncaya kadar ürünlerin her türlü hasar ve kayba karşı "
     "sigortalanmasından sorumludur. Sigorta poliçesinin bir örneği talep halinde Alıcı'ya "
     "iletilir."),
    ("FİKRİ MÜLKİYET", "Sözleşme kapsamında Satıcı tarafından Alıcı'ya özel olarak geliştirilen "
     "tasarım, doküman, yazılım ve benzeri her türlü çalışmanın fikri mülkiyet hakları Alıcı'ya "
     "aittir. Satıcı, teslim ettiği ürünlerin üçüncü kişilerin fikri ve sınai mülkiyet haklarını "
     "ihlal etmediğini taahhüt eder; aksi halde doğacak her türlü talepten Alıcı'yı ari tutar."),
    ("ALT YÜKLENİCİ", "Satıcı, sözleşme konusu işlerin tamamını veya bir kısmını Alıcı'nın önceden "
     "yazılı onayı olmaksızın alt yüklenicilere yaptıramaz. Onay verilmesi halinde dahi Satıcı, "
     "alt yüklenicinin fiil ve ihmallerinden Alıcı'ya karşı müteselsilen sorumludur."),
    ("İŞ SAĞLIĞI, GÜVENLİĞİ VE ÇEVRE", "Satıcı, Alıcı'nın tesislerinde yürüteceği faaliyetlerde "
     "6331 sayılı İş Sağlığı ve Güvenliği Kanunu ile çevre mevzuatına ve Alıcı'nın iç "
     "talimatlarına uymakla yükümlüdür. Bu yükümlülüklere aykırılıktan doğan idari para cezaları "
     "ve tazminatlar Satıcı tarafından karşılanır. Alıcı, aykırılığın tespiti halinde işi "
     "durdurma hakkına sahiptir."),
    ("TEBLİGAT", "Tarafların işbu sözleşmede yazılı adresleri yasal tebligat adresleridir. Adres "
     "değişiklikleri yazılı olarak bildirilmedikçe bu adreslere yapılan tebligatlar geçerli "
     "sayılır. Taraflar ayrıca kayıtlı elektronik posta adresleri üzerinden de tebligat "
     "yapılabileceğini kabul eder."),
    ("DAMGA VERGİSİ VE MASRAFLAR", "İşbu sözleşmeden doğan damga vergisi taraflarca eşit olarak "
     "karşılanır. Sözleşmenin ifasıyla ilgili diğer vergi, resim ve harçlar ilgili mevzuat "
     "uyarınca yükümlüsü olan tarafça ödenir."),
    ("SÖZLEŞMENİN EKLERİ", "Teknik şartname (EK-1), fiyat listesi (EK-2) ve teslimat planı "
     "(EK-3) işbu sözleşmenin ayrılmaz parçasıdır. Sözleşme hükümleri ile ekler arasında "
     "çelişki bulunması halinde sözleşme hükümleri esas alınır."),
    ("UYUŞMAZLIKLARIN ÇÖZÜMÜ", "İşbu sözleşmeden doğabilecek uyuşmazlıkların çözümünde Türk "
     "hukuku uygulanır ve İstanbul (Çağlayan) Mahkemeleri ile İcra Daireleri yetkilidir."),
]  # fmt: skip

PT_PER_MM = 72 / 25.4


def find_font() -> tuple[Path, Path | None]:
    if env := os.environ.get("MOCK491_FONT"):
        return Path(env), None
    for regular, bold in FONT_CANDIDATES:
        if regular.exists():
            return regular, bold if bold.exists() else None
    raise RuntimeError(
        "No Unicode TTF font found. Put DejaVuSans.ttf in assets/fonts/ or set MOCK491_FONT."
    )


def tr_number(x: float, decimals: int = 2) -> str:
    """250000.5 -> '250.000,50' (Turkish thousands/decimal separators)."""
    return f"{x:,.{decimals}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def tr_percent(rate: float) -> str:
    return "%" + f"{rate * 100:g}".replace(".", ",")


@dataclass
class ContractSpec:
    contract_no: str
    buyer_name: str
    supplier_text: str
    vkn_text: str | None
    amount: float
    currency: str
    penalty_rate_daily: float | None
    signed_on: date
    duration_months: int
    filler: list[tuple[str, str]]
    price_after: int  # how many filler clauses come before the price clause


@dataclass
class ClauseLocation:
    page_number: int
    bounding_box: tuple[float, float, float, float]  # points, top-left origin


def render_contract(spec: ContractSpec, path: Path) -> ClauseLocation:
    """Write the contract PDF; return where the price/penalty clause landed."""
    regular, bold = find_font()
    pdf = FPDF(unit="pt", format="A4")
    pdf.set_creation_date(datetime(2026, 1, 1, tzinfo=UTC))
    pdf.set_margins(56, 56, 56)
    pdf.set_auto_page_break(True, margin=56)
    pdf.add_font("body", "", regular)
    pdf.add_font("body", "B", bold or regular)
    pdf.add_page()

    def heading(text: str) -> None:
        pdf.set_font("body", "B", 11)
        pdf.ln(8)
        pdf.multi_cell(0, 15, text, new_x="LMARGIN", new_y="NEXT")

    def para(text: str) -> None:
        pdf.set_font("body", "", 11)
        pdf.multi_cell(0, 16, text, new_x="LMARGIN", new_y="NEXT", align="J")

    pdf.set_font("body", "B", 15)
    pdf.cell(0, 22, "TEDARİK SÖZLEŞMESİ", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("body", "", 10)
    pdf.cell(0, 14, f"Sözleşme No: {spec.contract_no}    Tarih: {spec.signed_on:%d.%m.%Y}",
             align="C", new_x="LMARGIN", new_y="NEXT")  # fmt: skip

    vkn_part = f" (VKN: {spec.vkn_text})" if spec.vkn_text else ""
    clauses = [
        ("TARAFLAR", f"İşbu sözleşme, bir tarafta {spec.buyer_name} (\"Alıcı\") ile diğer tarafta "
         f"{spec.supplier_text}{vkn_part} (\"Satıcı\") arasında aşağıdaki şartlarla "
         "imzalanmıştır."),
        ("KONU", "İşbu sözleşmenin konusu, Satıcı tarafından Alıcı'ya ekli teknik şartnamede "
         "belirtilen mal ve hizmetlerin tedarik edilmesine ilişkin usul ve esasların "
         f"belirlenmesidir. Sözleşme süresi imza tarihinden itibaren {spec.duration_months} "
         "aydır."),
        *spec.filler,
    ]  # fmt: skip
    price = (f"İşbu sözleşmenin toplam bedeli {tr_number(spec.amount)} {spec.currency}'dir. "
             "Bedel, faturanın tebliğinden itibaren altmış (60) gün içinde ödenir.")  # fmt: skip
    if spec.penalty_rate_daily is not None:
        price += (" Satıcının teslim yükümlülüğünü gecikmeli ifa etmesi halinde, geciken her gün "
                  f"için sözleşme bedelinin günlük {tr_percent(spec.penalty_rate_daily)} oranında "
                  "cezai şart uygulanır.")  # fmt: skip
    # The price clause sits after a random number of filler clauses, so its page varies.
    clauses.insert(2 + spec.price_after, ("BEDEL VE CEZAİ ŞART", price))

    location = None
    for n, (title, body) in enumerate(clauses, start=1):
        if pdf.will_page_break(70):  # keep each heading with the start of its clause
            pdf.add_page()
        heading(f"MADDE {n} - {title}")
        if title == "BEDEL VE CEZAİ ŞART":
            page, top = pdf.page_no(), pdf.get_y()
            para(body)
            location = ClauseLocation(page, (pdf.l_margin, top, pdf.w - pdf.r_margin, pdf.get_y()))
        else:
            para(body)

    pdf.ln(30)
    pdf.set_font("body", "B", 10.5)
    pdf.cell(pdf.epw / 2, 14, "ALICI")
    pdf.cell(pdf.epw / 2, 14, "SATICI", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("body", "", 10)
    pdf.cell(pdf.epw / 2, 14, spec.buyer_name)
    pdf.cell(pdf.epw / 2, 14, spec.supplier_text, new_x="LMARGIN", new_y="NEXT")

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    assert location is not None
    return location


@dataclass
class InvoiceRow:
    invoice_no: str
    issued_on: date
    supplier_text: str
    vkn_text: str | None
    amount: float
    currency: str
    paid: bool


# Two subsidiaries never agree on a spreadsheet layout (report section 9, question 3).
INVOICE_DIALECTS = {
    "A": ["Fatura No", "Fatura Tarihi", "Firma", "Vergi_No", "Tutar", "Para Birimi", "Durum"],
    "B": ["FATURA_NO", "TARIH", "TEDARIKCI_UNVANI", "VKN", "TUTAR", "ODEME_DURUMU"],
}

CURRENCY_LABEL = {"TRY": "TL", "USD": "USD", "EUR": "EUR"}


def render_invoices(rows: list[InvoiceRow], dialect: str, path: Path) -> None:
    """Write one workbook. Data rows start at Excel row 2."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturalar"
    ws.append(INVOICE_DIALECTS[dialect])
    for r in rows:
        if dialect == "A":
            ws.append([r.invoice_no, r.issued_on, r.supplier_text, r.vkn_text, r.amount,
                       r.currency, "Ödendi" if r.paid else "Ödenmedi"])  # fmt: skip
            ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY"
            ws.cell(ws.max_row, 5).number_format = "#,##0.00"
        else:
            ws.append([r.invoice_no, f"{r.issued_on:%d.%m.%Y}", r.supplier_text, r.vkn_text or "",
                       f"{tr_number(r.amount)} {CURRENCY_LABEL[r.currency]}",
                       "ÖDENDİ" if r.paid else "ÖDENMEDİ"])  # fmt: skip
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = (
            max(len(str(c.value or "")) for c in col) + 2
        )
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def pick_filler(rng: random.Random) -> list[tuple[str, str]]:
    return rng.sample(FILLER_CLAUSES, rng.randint(4, len(FILLER_CLAUSES)))
