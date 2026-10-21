"""Open the original bronze file at the place an extraction came from ("show me the source")."""

from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl import load_workbook
from PIL import Image

HIGHLIGHT_PAD = 4  # pt around the stored bounding box


def pdf_page(path: Path, page_number: int, bbox: tuple[float, float, float, float] | None,
             resolution: int = 110, crop_margin: float | None = None) -> Image.Image:  # fmt: skip
    """Render one page (1-based) with the extracted clause outlined.

    `bbox` is (x0, top, x1, bottom) in pt with a top-left origin, as stored in lineage,
    which is the coordinate system pdfplumber draws in. With `crop_margin` (pt), only a
    full-width band around the clause is returned: the "evidence" snippet.
    """
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[page_number - 1]
        im = page.to_image(resolution=resolution)
        if bbox:
            x0, top, x1, bottom = bbox
            box = (x0 - HIGHLIGHT_PAD, top - HIGHLIGHT_PAD, x1 + HIGHLIGHT_PAD,
                   bottom + HIGHLIGHT_PAD)  # fmt: skip
            im.draw_rect(box, fill=(255, 200, 0, 60), stroke=(220, 120, 0), stroke_width=3)
        image = im.annotated.convert("RGB")
        if bbox and crop_margin is not None:
            scale = resolution / 72
            upper = max(0.0, bbox[1] - crop_margin) * scale
            lower = min(float(page.height), bbox[3] + crop_margin) * scale
            image = image.crop((0, int(upper), image.width, int(lower)))
        return image


def sheet_rows(path: Path) -> pd.DataFrame:
    """The active sheet as written, indexed by Excel row number (header is row 1)."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    df = pd.DataFrame(list(rows), columns=header)
    df.index = range(2, len(df) + 2)
    df.index.name = "row"
    return df
