import re
from collections import Counter

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from mock491 import mockgen, vkn
from mock491.mockgen import render


def squash(s: str) -> str:
    return re.sub(r"\s+", " ", s)


def test_vkn_checksum():
    full = "123456789" + str(vkn.check_digit("123456789"))
    assert vkn.is_valid(full)
    assert not vkn.is_valid(full[:9] + str((int(full[9]) + 1) % 10))


def test_counts_and_files(dataset):
    out, truth = dataset
    assert len(truth.entities) == 13  # 10 suppliers + 3 look-alikes
    assert Counter(r.doc_type for r in truth.records) == {"contract": 12, "invoice": 32}
    assert all((out / r.file).exists() for r in truth.records)
    assert all(vkn.is_valid(e.vkn) for e in truth.entities)
    assert mockgen.load_truth(out) == truth
    assert {r.entity_id for r in truth.records} == {e.entity_id for e in truth.entities}


def test_lookalikes_are_distinct_companies(dataset):
    _, truth = dataset
    by_id = {e.entity_id: e for e in truth.entities}
    for e in truth.entities:
        if e.sibling_of:
            other = by_id[e.sibling_of]
            assert e.vkn != other.vkn
            assert e.canonical_name.split()[0] == other.canonical_name.split()[0]


def test_contract_pdf_matches_truth(dataset):
    out, truth = dataset
    for r in (r for r in truth.records if r.doc_type == "contract"):
        page = squash(PdfReader(out / r.file).pages[r.page_number - 1].extract_text())
        assert f"{render.tr_number(r.amount)} {r.currency}" in page, r.record_id
        assert squash(r.name_as_written) in squash(PdfReader(out / r.file).pages[0].extract_text())


def test_invoice_rows_match_truth(dataset):
    out, truth = dataset
    for r in (r for r in truth.records if r.doc_type == "invoice"):
        ws = load_workbook(out / r.file).active
        row = [c.value for c in ws[r.row_number]]
        assert row[0] == r.record_id and row[2] == r.name_as_written


def test_same_seed_same_truth(dataset, tmp_path):
    _, truth = dataset
    again = mockgen.generate(tmp_path / "again", seed=7, n_suppliers=10,
                             contracts_per_tenant=3, invoices_per_tenant=8)  # fmt: skip
    assert again == truth


def test_refuses_to_wipe_unrelated_folder(tmp_path):
    (tmp_path / "keep.txt").write_text("important")
    with pytest.raises(FileExistsError):
        mockgen.generate(tmp_path)
